import json
from pathlib import Path
import random
import tempfile
import unittest
from openpyxl import load_workbook
from matching.models import Config, Module
from matching.algorithm import solve, cost, violations
from matching.data import prepare, tables_in_file, unique_rows
from matching.cli import run
import pandas as pd

ROOT = Path(__file__).resolve().parent.parent


def brute_force(modules, config, remaining=None):
    remaining = list(range(len(modules))) if remaining is None else remaining
    if not remaining:
        return (0, 0)
    first, *tail = remaining
    best = brute_force(modules, config, tail)
    for other in tail:
        if violations(modules[first], modules[other], config):
            continue
        count, total = brute_force(modules, config, [i for i in tail if i != other])
        candidate = (count + 1, total + round(cost(modules[first], modules[other], config) * 1_000_000))
        if candidate[0] > best[0] or (candidate[0] == best[0] and candidate[1] < best[1]):
            best = candidate
    return best


class AlgorithmTests(unittest.TestCase):
    def test_default_score_is_enabled(self):
        config = Config()
        self.assertTrue(config.use_score)
        self.assertEqual(config.score_source, "block-average")
        self.assertTrue(json.loads((ROOT / "config.json").read_text(encoding="utf-8"))["use_score"])
        modules = [Module("a", 0, 10, 10, 0, score=40), Module("b", 0, 10, 10, 0, score=47)]
        self.assertEqual(solve(modules, config)[0], [])
        self.assertEqual(len(solve(modules, Config(use_score=False))[0]), 1)

    def test_exact_against_enumeration_in_both_modes_and_ab(self):
        rng = random.Random(42)
        for mode, grouped in [("no-imu", False), ("mixed-imu", False), ("no-imu", True), ("mixed-imu", True)]:
            for _ in range(25):
                config = Config(mode=mode, use_score=False, weights={"aec": .2, "rg": .4, "bg": .4, "score": 0})
                modules = [Module(str(i), rng.randint(0, 3), rng.randint(0, 20), rng.randint(0, 20),
                                  int(i < 4) if mode == "mixed-imu" else 0,
                                  "A" if i < 4 else "B" if grouped else "pool") for i in range(8)]
                # For ungrouped mode every node belongs to the shared pool.
                if not grouped:
                    modules = [Module(m.sn, m.aec, m.rg, m.bg, m.imu) for m in modules]
                pairs, unpaired, _ = solve(modules, config)
                actual = (len(pairs), sum(round(cost(modules[i], modules[j], config) * 1_000_000) for i, j in pairs))
                self.assertEqual(actual, brute_force(modules, config))
                self.assertEqual(len({i for pair in pairs for i in pair}), 2 * len(pairs))
                self.assertEqual(2 * len(pairs) + len(unpaired), len(modules))

    def test_boundary_score_and_exposure_window(self):
        modules = [Module("a", 1, 10, 10, 0, score=0), Module("b", 2, 20, 20, 0, score=6.8)]
        self.assertEqual(len(solve(modules, Config(use_score=True))[0]), 1)
        self.assertEqual(len(solve(modules, Config(use_score=True, score_limit=6.7))[0]), 0)
        modules[1] = Module("b", 2.001, 10, 10, 0)
        self.assertEqual(solve(modules, Config(use_score=False))[0], [])

    def test_cardinality_beats_one_perfect_pair(self):
        modules = [Module(str(i), 0, rg, 0, 0) for i, rg in enumerate([0, 10, 10, 20])]
        pairs, _, _ = solve(modules, Config(use_score=False))
        self.assertEqual(len(pairs), 2)

    def test_config_rejects_nonfinite_and_invalid_limits(self):
        for value in (0, -1, float("nan"), float("inf")):
            with self.assertRaises(ValueError):
                Config(aec_limit=value)
        with self.assertRaises(ValueError):
            Config(use_score=False, weights={"aec": 0, "rg": 0, "bg": 0, "score": 1})


class InputTests(unittest.TestCase):
    def make_table(self, directory, name, text):
        path = Path(directory) / name
        path.write_text(text, encoding="utf-8")
        return path

    def test_missing_numeric_is_invalid_not_zero(self):
        with tempfile.TemporaryDirectory() as tmp:
            p = self.make_table(tmp, "data.csv", "SN,AEC,R/G,B/G\na,1,10,10\nb,,10,10\nc,inf,10,10\n")
            modules, invalid, _ = prepare([(p, "pool", 0)], Config(use_score=False))
            self.assertEqual([m.sn for m in modules], ["a"])
            self.assertEqual(len(invalid), 2)

    def test_imu_is_explicit_and_conflicts_not_ignored(self):
        with tempfile.TemporaryDirectory() as tmp:
            p = self.make_table(tmp, "data.csv", "SN,AEC,R/G,B/G,IMU_Flag\na,1,10,10,\nb,1,10,10,1\nc,1,10,10,0.5\n")
            mods, invalid, _ = prepare([(p, "pool", None)], Config(mode="mixed-imu", use_score=False))
            self.assertEqual([m.sn for m in mods], ["b"])
            self.assertEqual(len(invalid), 2)
            mods, invalid, _ = prepare([(p, "A", 1)], Config(mode="mixed-imu", use_score=False))
            self.assertEqual(len(mods), 2)
            self.assertEqual(len(invalid), 1)
            with self.assertRaises(ValueError):
                run(input_path=p, output=tmp, config=Config(mode="mixed-imu", use_score=False))

    def test_conflicting_retests_are_rejected(self):
        identical = pd.DataFrame({"SN.": ["a", "a"], "AEC": ["1", "1.0"]})
        self.assertEqual(len(unique_rows(identical, "AEC")), 1)
        with self.assertRaises(ValueError):
            unique_rows(pd.DataFrame({"SN.": ["a", "a"], "AEC": [1, 2]}), "AEC")

    def test_blank_rows_do_not_truncate_table(self):
        with tempfile.TemporaryDirectory() as tmp:
            p = self.make_table(tmp, "data.csv", "SN,AEC,R/G,B/G\na,1,10,10\n\nb,1,10,10\n")
            modules, _, _ = prepare([(p, "pool", 0)], Config(use_score=False))
            self.assertEqual(len(modules), 2)

    def test_cross_group_duplicate_sn_is_rejected(self):
        with tempfile.TemporaryDirectory() as tmp:
            p = self.make_table(tmp, "data.csv", "SN,AEC,R/G,B/G\na,1,10,10\n")
            with self.assertRaises(ValueError):
                prepare([(p, "A", 0), (p, "B", 0)], Config(use_score=False))

    def test_partial_sn_is_retained_as_invalid(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.make_table(tmp, "mtf.csv", "SN,Block3_H,Block4_H\na,40,40\nb,40,40\n")
            self.make_table(tmp, "aec.csv", "SN,AEC\na,1\nb,1\n")
            self.make_table(tmp, "awb.csv", "SN,R/G,B/G\na,10,10\n")
            modules, invalid, _ = prepare([(tmp, "pool", 0)], Config(use_score=False))
            self.assertEqual(len(modules), 1)
            self.assertEqual(invalid[0]["SN"], "b")

    def test_score_defaults_to_block_average_even_when_filter_disabled(self):
        with tempfile.TemporaryDirectory() as tmp:
            p = self.make_table(tmp, "data.csv", "SN,AEC,R/G,B/G,Block3_H,Block4_H\na,1,10,10,40,42\n")
            modules, invalid, _ = prepare([(p, "pool", 0)], Config(use_score=True, score_source="column"))
            self.assertEqual(len(invalid), 1)
            modules, invalid, _ = prepare([(p, "pool", 0)], Config(use_score=True))
            self.assertEqual(modules[0].score, 41)
            modules, invalid, _ = prepare([(p, "pool", 0)], Config(use_score=False))
            self.assertEqual(modules[0].score, 41)
            self.assertFalse(invalid)
            p.write_text("SN,AEC,R/G,B/G,Block3_H,Block4_H\na,1,10,10,40,\n", encoding="utf-8")
            modules, invalid, _ = prepare([(p, "pool", 0)], Config(use_score=True))
            self.assertFalse(modules)
            self.assertEqual(len(invalid), 1)


class IntegrationTests(unittest.TestCase):
    @unittest.skipUnless((ROOT / "input/GC6603_AEC_Data_20261009.xls").exists() and (ROOT / "input/0922/GC6603_AEC_Data_20260922.xls").exists(), "本地产线数据未提供；公开仓库使用合成示例验证")
    def test_current_batch_and_historical_batch(self):
        modules, invalid, _ = prepare([(ROOT / "input", "pool", 0)], Config(use_score=False))
        self.assertEqual((len(modules), len(invalid)), (10, 0))
        pairs, unpaired, _ = solve(modules, Config(use_score=False))
        self.assertEqual((len(pairs), len(unpaired)), (5, 0))
        self.assertEqual(sum(abs(modules[i].rg-modules[j].rg) + abs(modules[i].bg-modules[j].bg) for i,j in pairs), 21)
        modules, invalid, _ = prepare([(ROOT / "input/0922", "pool", 0)], Config(use_score=False))
        self.assertTrue(modules)
        self.assertFalse(invalid)

    @unittest.skipUnless((ROOT / "260604/带IMU_10PCS_20260602.xlsx").exists() and (ROOT / "260604/不带IMU 10PCS_20260602.xlsx").exists(), "本地历史工作簿未提供")
    def test_actual_mixed_imu_workbooks(self):
        sources = [(ROOT / "260604/带IMU_10PCS_20260602.xlsx", "A", 1),
                   (ROOT / "260604/不带IMU 10PCS_20260602.xlsx", "B", 0)]
        config = Config(mode="mixed-imu", use_score=False)
        modules, invalid, _ = prepare(sources, config)
        self.assertEqual((len(modules), len(invalid)), (20, 0))
        pairs, _, _ = solve(modules, config)
        self.assertTrue(pairs)
        for i, j in pairs:
            self.assertEqual((modules[i].imu, modules[j].imu), (1, 0))
            self.assertEqual((modules[i].group, modules[j].group), ("A", "B"))
            self.assertFalse(violations(modules[i], modules[j], config))

    def test_one_command_report_and_no_overwrite(self):
        with tempfile.TemporaryDirectory() as tmp:
            first = run(input_path=ROOT / "examples/noimu", output=tmp)
            second = run(input_path=ROOT / "examples/noimu", output=tmp)
            self.assertNotEqual(first, second)
            book = load_workbook(first, data_only=False)
            self.assertEqual(book.sheetnames, ["Pairs", "Unpaired", "Invalid", "Data", "Config"])
            self.assertEqual(book["Pairs"].max_row, 6)
            headers = [cell.value for cell in book["Pairs"][1]]
            self.assertEqual(headers.count("IMU_Flag_A"), 1)
            self.assertEqual(headers.count("IMU_Flag_B"), 1)
            self.assertNotIn("IMU_A", headers)
            self.assertNotIn("IMU_B", headers)
            self.assertEqual(book["Data"].max_row, 11)
            self.assertEqual(book["Invalid"].max_row, 1)
            summary = json.loads((first.parent / "summary.json").read_text(encoding="utf-8"))
            self.assertEqual(summary["配对数"], 5)

    def test_public_mixed_example(self):
        config = Config(mode="mixed-imu", use_score=True)
        sources = [(ROOT / "examples/mixed/A.csv", "A", 1), (ROOT / "examples/mixed/B.csv", "B", 0)]
        modules, invalid, _ = prepare(sources, config)
        pairs, unpaired, _ = solve(modules, config)
        self.assertEqual((len(pairs), len(unpaired), len(invalid)), (3, 0, 0))

    def test_input_identifiers_are_not_excel_formulas(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "data.csv"
            source.write_text("SN,AEC,R/G,B/G,Block3_H,Block4_H\n=1+1,1,10,10,40,40\nb,1,10,10,40,40\n", encoding="utf-8")
            target = run(input_path=source, output=Path(tmp) / "results")
            book = load_workbook(target, data_only=False)
            cells = [cell for row in book["Pairs"] for cell in row if cell.value == "=1+1"]
            self.assertTrue(cells)
            self.assertTrue(all(cell.data_type == "s" for cell in cells))


if __name__ == "__main__":
    unittest.main()
