"""Regression coverage for audit trails, workflow boundaries and atomic exports."""
from dataclasses import replace
from io import StringIO
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import pandas as pd
from openpyxl import load_workbook
from matching.data import prepare, tables_in_file
from matching.models import Config
from matching.errors import InputError
from matching.service import execute
from matching.runtime import read_preferences, save_preferences

ROOT = Path(__file__).resolve().parent.parent


class WorkflowTests(unittest.TestCase):
    def test_missing_sn_is_audited_with_physical_csv_line(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / 'batch.csv'
            source.write_text('SN,AEC,R/G,B/G\na,1,10,10\n\n,1,10,10\nNA,1,10,10\n', encoding='utf-8')
            modules, invalid, _ = prepare([(source, 'pool', 0)], Config(use_score=False))
            self.assertEqual([m.sn for m in modules], ['NA', 'a'])
            self.assertEqual(len(invalid), 1)
            self.assertEqual(invalid[0]['错误代码'], 'missing_sn')
            self.assertIn('第4行', invalid[0]['来源'])
            self.assertIn('第5行', modules[0].sources[0])

    def test_all_missing_sn_is_not_silently_skipped(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / 'batch.csv'
            source.write_text('SN,AEC,R/G,B/G\n,1,10,10\n', encoding='utf-8')
            result = execute(source, output=Path(tmp) / 'out', config=Config(use_score=False))
            self.assertEqual(result.summary['有效模组'], 0)
            book = load_workbook(result.target)
            self.assertEqual(book['Invalid'].max_row, 2)
            book.close()

    def test_excel_dedup_retains_sheet_rows_and_blank_rows(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / 'batch.xlsx'
            pd.DataFrame({'SN': ['a', None, 'a'], 'AEC': [1, None, 1], 'R/G': [10, None, 10], 'B/G': [10, None, 10]}).to_excel(source, sheet_name='光学测试', index=False)
            modules, invalid, _ = prepare([(source, 'pool', 0)], Config(use_score=False))
            self.assertFalse(invalid)
            self.assertEqual(len(modules), 1)
            self.assertEqual(len(modules[0].sources), 2)
            self.assertIn('光学测试 / 第2行', modules[0].sources[0])
            self.assertIn('光学测试 / 第4行', modules[0].sources[1])

    def test_conflict_error_lists_fields_and_both_locations(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / 'batch.csv'
            source.write_text('SN,AEC,R/G,B/G\na,1,10,10\na,2,10,10\n', encoding='utf-8')
            with self.assertRaises(InputError) as caught:
                prepare([(source, 'pool', 0)], Config(use_score=False))
            self.assertEqual(caught.exception.code, 'conflicting_retest')
            for detail in ['AEC', '第2行', '第3行', str(source.resolve())]:
                self.assertIn(detail, str(caught.exception))

    def test_split_measurements_retain_all_three_sources(self):
        with tempfile.TemporaryDirectory() as tmp:
            for name, text in [('mtf', 'SN,Block3_H,Block4_H\na,40,42\n'), ('aec', 'SN,AEC\na,1\n'), ('awb', 'SN,R/G,B/G\na,10,10\n')]:
                (Path(tmp) / (name + '.csv')).write_text(text, encoding='utf-8')
            modules, invalid, _ = prepare([(tmp, 'pool', 0)], Config())
            self.assertFalse(invalid)
            self.assertEqual(len(modules[0].sources), 3)

    def test_service_has_no_console_side_effects_and_emits_stages(self):
        with tempfile.TemporaryDirectory() as tmp:
            events = []
            stream = StringIO()
            with patch('sys.stdout', stream):
                result = execute(ROOT / 'examples/noimu', output=tmp, progress=lambda stage, message: events.append(stage))
            self.assertEqual(stream.getvalue(), '')
            self.assertEqual(set(events), {'read', 'match', 'export', 'done'})
            self.assertEqual(result.summary['配对数'], 5)
            self.assertEqual(result.summary['状态'], 'complete')
            self.assertIn('读取校验秒', result.summary['阶段耗时秒'])
            book = load_workbook(result.target)
            headers = [c.value for c in book['Data'][1]]
            self.assertIn('来源', headers)
            book.close()

    def test_mode_validation_uses_codes_not_display_text(self):
        from matching.models import Module
        invalid = [{'SN': 'bad', '来源': 'test', '原因': 'Translated text', '错误代码': 'imu_conflict'}]
        with patch('matching.service.prepare', return_value=([Module('a', 1, 1, 1, 0)], invalid, [])):
            with self.assertRaises(InputError) as caught:
                execute('unused')
        self.assertEqual(caught.exception.code, 'imu_conflict')

    def test_failed_summary_write_does_not_publish_partial_result(self):
        with tempfile.TemporaryDirectory() as tmp:
            with patch('matching.report.Path.write_text', side_effect=OSError('disk full')):
                with self.assertRaises(OSError):
                    execute(ROOT / 'examples/noimu', output=tmp)
            self.assertEqual(list(Path(tmp).iterdir()), [])

    def test_preferences_roundtrip_and_corrupt_file_fallback(self):
        with tempfile.TemporaryDirectory() as tmp:
            target = Path(tmp) / 'settings.json'
            data = {'input': '中文目录', 'config': {'use_score': False}}
            save_preferences(data, target)
            self.assertEqual(read_preferences(target), data)
            target.write_text('{bad', encoding='utf-8')
            self.assertEqual(read_preferences(target), {})
            self.assertEqual([p.name for p in Path(tmp).iterdir()], ['settings.json'])

    def test_csv_multiline_values_keep_original_row_numbers(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / 'batch.csv'
            source.write_text('SN,AEC,R/G,B/G\n"a\ncontinued",1,10,10\nb,1,10,10\n', encoding='utf-8')
            tables = tables_in_file(source, [])
            self.assertIn('第4行', tables[0][1].iloc[1]['__location'])


class DiagnosticTests(unittest.TestCase):
    def test_index_matches_full_scan_across_modes_weights_and_ties(self):
        import random
        import networkx as nx
        from matching.algorithm import DiagnosticIndex, explain_unpaired
        from matching.models import Module
        rng = random.Random(2026)
        for mode in ['no-imu', 'mixed-imu']:
            for grouped in [False, True]:
                for use_score in [False, True]:
                    for weights in [dict(aec=0, rg=.5, bg=.5, score=0), dict(aec=.2, rg=.2, bg=.2, score=.4), dict(aec=1, rg=0, bg=0, score=0)]:
                        config = Config(mode=mode, use_score=use_score, weights=weights)
                        modules = [Module(f'SN{i:03}', rng.randrange(10), rng.randrange(15), rng.randrange(15), i % 2,
                                          ('A' if i % 2 else 'B') if grouped else 'pool', score=rng.randrange(15)) for i in range(50)]
                        graph = nx.empty_graph(len(modules))
                        index = DiagnosticIndex(modules, config)
                        for i in range(len(modules)):
                            self.assertEqual(explain_unpaired(i, modules, graph, config, index), explain_unpaired(i, modules, graph, config))

    def test_index_handles_missing_scores_and_no_opposite_group(self):
        from matching.algorithm import DiagnosticIndex
        from matching.models import Module
        config = Config(weights=dict(aec=0, rg=0, bg=0, score=1))
        modules = [Module('b', 1, 1, 1, 0, score=None), Module('a', 1, 1, 1, 0, score=1)]
        self.assertEqual(DiagnosticIndex(modules, config).nearest(0), 1)
        self.assertIsNone(DiagnosticIndex([replace(m, group='A') for m in modules], config).nearest(0))


if __name__ == '__main__':
    unittest.main()
