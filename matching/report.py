"""One self-contained result workbook per run; never overwrite history."""
from dataclasses import asdict
from datetime import datetime
from pathlib import Path
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from .algorithm import differences, cost, explain_unpaired


def module_row(module):
    return {"SN": module.sn, "Group": module.group, "IMU_Flag": module.imu,
            "AEC": module.aec, "R/G": module.rg, "B/G": module.bg,
            "Score": module.score, "Block3_H": module.block3, "Block4_H": module.block4}


def write_sheet(book, name, rows, empty_headers):
    sheet = book.create_sheet(name)
    headers = list(dict.fromkeys(key for row in rows for key in row)) if rows else empty_headers
    sheet.append(headers)
    for row in rows:
        # Strings from input files are identifiers, never executable Excel formulas.
        sheet.append([row.get(key) for key in headers])
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="24476B")
    for column in sheet.columns:
        length = max(len(str(cell.value or "")) for cell in column)
        sheet.column_dimensions[column[0].column_letter].width = min(65, max(14, length * 1.3 + 2))
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, float):
                cell.number_format = "0.0000"
    return sheet


def export(output, modules, pairs, unpaired, graph, invalid, notes, config, sources):
    output = Path(output)
    output.mkdir(parents=True, exist_ok=True)
    run = output / datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    run.mkdir(exist_ok=False)
    book = Workbook()
    book.remove(book.active)
    pair_rows = []
    for pair_id, (i, j) in enumerate(pairs, 1):
        a, b = modules[i], modules[j]
        pair_rows.append({"Pair_ID": pair_id, "SN_A": a.sn, "SN_B": b.sn,
                          "IMU_Flag_A": a.imu, "IMU_Flag_B": b.imu,
                          **{f"Δ{k}": v for k, v in differences(a, b).items()},
                          "归一化成本": cost(a, b, config),
                          **{f"{k}_A": v for k, v in module_row(a).items() if k not in {"SN", "IMU_Flag"}},
                          **{f"{k}_B": v for k, v in module_row(b).items() if k not in {"SN", "IMU_Flag"}}})
    write_sheet(book, "Pairs", pair_rows, ["Pair_ID", "SN_A", "SN_B", "IMU_Flag_A", "IMU_Flag_B", "Δaec", "Δrg", "Δbg", "Δscore", "归一化成本"])
    write_sheet(book, "Unpaired", [{**module_row(modules[i]), **explain_unpaired(i, modules, graph, config)} for i in unpaired],
                ["SN", "原因", "合格候选数", "最近候选SN"])
    write_sheet(book, "Invalid", invalid, ["SN", "Group", "来源", "原因"])
    write_sheet(book, "Data", [module_row(m) for m in modules], ["SN", "Group", "IMU_Flag", "AEC", "R/G", "B/G", "Score"])
    summary = {"模式": config.mode, "有效模组": len(modules), "配对数": len(pairs),
               "未配对数": len(unpaired), "无效模组": len(invalid), "合格边数": graph.number_of_edges(),
               "总归一化成本": sum(cost(modules[i], modules[j], config) for i, j in pairs),
               "成本精度": "每条边取整至1e-6，等价最优解可能有多个",
               "目标": "先最多配对，再最小总成本；严格阈值，不自动放宽",
               "输入": [{"path": str(Path(p).resolve()), "group": group, "assigned_imu": imu} for p, group, imu in sources],
               "配置": asdict(config), "读取说明": notes}
    config_rows = [{"项目": k, "值": json.dumps(v, ensure_ascii=False) if isinstance(v, (dict, list)) else v} for k, v in summary.items() if k != "读取说明"]
    config_rows.extend({"项目": "读取说明", "值": note} for note in notes)
    write_sheet(book, "Config", config_rows, ["项目", "值"])
    target = run / "匹配结果.xlsx"
    book.save(target)
    (run / "summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    return target
